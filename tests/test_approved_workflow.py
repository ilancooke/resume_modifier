from contextlib import ExitStack, redirect_stdout
from io import StringIO
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document

from approved_fixtures import make_inputs, make_plan, make_selection
from approved_tailoring import FeedbackDecision, ReserveBullet
from approved_workflow import publish_pair, run_approved_workflow
from config import AppConfig
from docx_editor import write_tailored_resume
from feedback_revision import collect_feedback_signals
from resume_reader import ResumeText
from test_feedback_revision import make_report


class ApprovedWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.library, self.template = make_inputs(self.root)
        self.plan = make_plan()
        self.initial = make_selection()
        self.job = self.root / "job.txt"
        self.job.write_text("Build production risk models with Python and AI evaluation.")
        self.config = AppConfig(self.template.path, self.root / "unused-narrative-history", self.job,
                                self.root / "outputs", self.root / "diagnostics", self.library.root)
        self.args = SimpleNamespace(resume=None, bullet_library=None, jd=None, model="gpt-6.1-terra",
                                    dry_run=False, show_diff=True, skip_feedback_revision=False)
        self.report = make_report(weak=True)
        self.signals = tuple(collect_feedback_signals(self.report))
        role = self.initial.roles[0].model_copy(update={"bullet_ids": (*self.initial.roles[0].bullet_ids, "ADS-OPS-01")})
        self.revised = self.initial.model_copy(update={
            "roles": (role, *self.initial.roles[1:]),
            "priority_order": (*self.initial.priority_order, "ADS-OPS-01"),
            "feedback_decisions": tuple(FeedbackDecision(signal_id=s.signal_id, status="addressed",
                rationale="Selected production monitoring and operational reliability evidence.", evidence_refs=("ADS-OPS-01",)) for s in self.signals),
        })

    @property
    def docx(self):
        return self.config.tailored_resumes_dir / "example/candidate_example_123.docx"

    @property
    def pdf(self):
        return self.docx.with_suffix(".pdf")

    def diagnostics(self):
        return next(self.config.resume_diagnostics_dir.glob("*.diagnostics.txt")).read_text()

    def archive(self):
        return next(self.config.resume_diagnostics_dir.glob("*.selection_runs/*"))

    @staticmethod
    def export(docx, pdf):
        # Exercise the real DOCX editor; only Word and the API are mocked here.
        pdf.write_text("\n".join(p.text for p in Document(docx).paragraphs))

    def services(self, *, responses=None, reports=None, layout=None):
        stack = ExitStack()
        self.addCleanup(stack.close)
        self.decide = stack.enter_context(patch("approved_workflow.decide_finbot", return_value=self.plan))
        self.select = stack.enter_context(patch("approved_workflow.request_selection", side_effect=responses or [self.initial, self.initial, self.revised, self.revised]))
        self.exports = stack.enter_context(patch("approved_workflow.export_docx_to_pdf", side_effect=self.export))
        self.layout = stack.enter_context(patch("approved_workflow.inspect_resume_pdf_layout", side_effect=layout or (lambda _: SimpleNamespace(has_overflow=False, page_count=2))))
        reports = iter(reports or [self.report, make_report(weak=False)])

        def evaluate(**kwargs):
            report = next(reports)
            if isinstance(report, BaseException):
                raise report
            path = kwargs["resume_path"]
            return ResumeText(path, path.read_text(), "pdf_text", 2, []), report

        self.match = stack.enter_context(patch("approved_workflow.validate_resume", side_effect=evaluate))

    def run_workflow(self):
        output = StringIO()
        with redirect_stdout(output):
            result = run_approved_workflow(args=self.args, config=self.config)
        return result, output.getvalue()

    def test_full_cycle_locks_finbot_and_reports_final_selection(self):
        self.services()
        result, output = self.run_workflow()
        self.assertEqual(result, 0)
        self.decide.assert_called_once()
        self.assertEqual(self.select.call_count, 4)
        self.assertTrue(all(call.kwargs["plan"] is self.plan for call in self.select.call_args_list))
        self.assertEqual(self.match.call_count, 2)
        self.assertIn("ADS-OPS-01", output)
        self.assertIn("Status: revised", self.diagnostics())
        self.assertIn("ADS-OPS-01", self.diagnostics())
        self.assertIn("Initial Resume Match:", self.diagnostics())
        self.assertIn("Tailored Resume Match:", self.diagnostics())
        self.assertEqual(json.loads((self.archive() / "locked_plan.json").read_text())["keep_finbot"], True)
        self.assertIn("production model pipelines", self.pdf.read_text())
        self.assertIn("SELECTED PROJECTS", self.pdf.read_text())
        self.assertEqual(self.docx.read_bytes(), (self.archive() / "candidate.docx").read_bytes())

    def test_removed_finbot_stays_removed_during_feedback_and_allows_extra_ds_bullets(self):
        self.plan = make_plan(False)
        self.initial = make_selection(10)
        decisions = tuple(FeedbackDecision(signal_id=s.signal_id, status="outside_editable_scope",
            rationale="Finbot was removed by the locked first-stage decision.", evidence_refs=()) for s in self.signals)
        unchanged = self.initial.model_copy(update={"feedback_decisions": decisions})
        self.services(responses=[self.initial, self.initial, unchanged])
        result, _ = self.run_workflow()
        self.assertEqual(result, 0)
        self.assertNotIn("Finbot", self.pdf.read_text())
        self.assertNotIn("SELECTED PROJECTS", self.pdf.read_text())
        self.assertIn("ADS-M9-01", self.diagnostics())
        self.decide.assert_called_once()
        self.assertEqual(self.select.call_count, 3)
        self.assertEqual(self.match.call_count, 1)

    def test_skills_and_retained_finbot_xml_and_package_parts_are_preserved(self):
        for keep in (True, False):
            with self.subTest(keep=keep):
                self.plan = make_plan(keep)
                # Directly verify both variants using the production writer.
                from approved_tailoring import assemble_resume
                output = self.root / f"variant-{keep}.docx"
                write_tailored_resume(template=self.template, tailoring=assemble_resume(self.initial, library=self.library, plan=self.plan),
                                      output_path=output, include_finbot=keep, role_bullet_limits=self.library.role_limits())
                source_paras = Document(self.template.path).paragraphs
                output_paras = Document(output).paragraphs
                for paragraph in source_paras:
                    if paragraph.text.startswith(("SKILLS", "Languages:", "LLM & AI Engineering:")) or (keep and paragraph.text in self.template.finbot_text().splitlines()):
                        actual = next(p for p in output_paras if p.text == paragraph.text)
                        self.assertEqual(actual._p.xml, paragraph._p.xml)
                with ZipFile(self.template.path) as source, ZipFile(output) as final:
                    for name in source.namelist():
                        if name != "word/document.xml":
                            self.assertEqual(source.read(name), final.read(name))

    def test_dry_run_makes_no_files_exports_or_match_calls(self):
        self.args.dry_run = True
        self.services()
        before = set(self.root.rglob("*"))
        result, output = self.run_workflow()
        self.assertEqual(result, 0)
        self.assertIn("Finbot: kept in full (locked)", output)
        self.assertEqual(set(self.root.rglob("*")), before)
        self.exports.assert_not_called()
        self.match.assert_not_called()
        self.assertEqual(self.select.call_count, 2)

    def test_no_feedback_and_explicit_skip_do_not_request_revision(self):
        for skip in (False, True):
            with self.subTest(skip=skip):
                self.args.skip_feedback_revision = skip
                self.services(reports=[make_report(weak=False)])
                result, _ = self.run_workflow()
                self.assertEqual(result, 0)
                self.assertEqual(self.select.call_count, 2)
                self.assertEqual(self.match.call_count, 1)

    def test_page_fitting_removes_approved_bullets_and_preserves_locked_content(self):
        self.args.skip_feedback_revision = True
        self.services(layout=[SimpleNamespace(has_overflow=True, page_count=3), SimpleNamespace(has_overflow=False, page_count=2)])
        result, _ = self.run_workflow()
        self.assertEqual(result, 0)
        self.assertEqual(self.exports.call_count, 2)
        outcome = json.loads((self.archive() / "outcome.json").read_text())
        self.assertNotIn("PC-01", outcome["final_selection"]["priority_order"])
        self.assertIn("Removed PC-01", self.diagnostics())
        self.assertIn("Evaluated AI extraction", self.pdf.read_text())

    def test_fills_extra_space_from_reviewed_reserves_without_extra_model_calls(self):
        self.args.skip_feedback_revision = True
        reserves = tuple(ReserveBullet(bullet_id=key, rationale="Adds relevant production ML evidence.",
                         second_bullet_rationale="") for key in ("ADS-M3-01", "ADS-M4-01"))
        selection = self.initial.model_copy(update={"reserve_bullets": reserves})
        self.services(responses=[selection, selection])
        result, output = self.run_workflow()
        self.assertEqual(result, 0)
        self.assertIn("ADS-M3-01", output)
        self.assertIn("ADS-M4-01", output)
        self.assertEqual(self.select.call_count, 2)
        self.assertEqual(self.exports.call_count, 3)
        self.assertIn("Added reserve ADS-M4-01", self.diagnostics())
        self.assertIn("SELECTED PROJECTS", self.pdf.read_text())

    def test_overflowing_reserve_keeps_last_valid_files_and_tries_next(self):
        self.args.skip_feedback_revision = True
        reserves = tuple(ReserveBullet(bullet_id=key, rationale="Relevant additional evidence.",
                         second_bullet_rationale="") for key in ("ADS-M3-01", "ADS-M4-01"))
        selection = self.initial.model_copy(update={"reserve_bullets": reserves})
        self.services(responses=[selection, selection], layout=[SimpleNamespace(has_overflow=False, page_count=2),
            SimpleNamespace(has_overflow=True, page_count=3), SimpleNamespace(has_overflow=False, page_count=2)])
        result, output = self.run_workflow()
        self.assertEqual(result, 0)
        self.assertNotIn("ADS-M3-01", output)
        self.assertIn("ADS-M4-01", output)
        self.assertIn("Deferred reserve ADS-M3-01", self.diagnostics())
        self.assertEqual(self.docx.read_bytes(), (self.archive() / "initial.docx").read_bytes())

    def test_optional_reserve_export_failure_retains_valid_draft(self):
        self.args.skip_feedback_revision = True
        reserve = ReserveBullet(bullet_id="ADS-M3-01", rationale="Relevant additional evidence.", second_bullet_rationale="")
        selection = self.initial.model_copy(update={"reserve_bullets": (reserve,)})
        self.services(responses=[selection, selection])
        calls = 0

        def export(docx, pdf):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("Word trial unavailable")
            self.export(docx, pdf)

        self.exports.side_effect = export
        result, output = self.run_workflow()
        self.assertEqual(result, 0)
        self.assertNotIn("ADS-M3-01", output)
        self.assertIn("retained the last valid two-page draft", self.diagnostics())
        self.assertEqual(self.docx.read_bytes(), (self.archive() / "initial.docx").read_bytes())

    def test_feedback_failures_keep_initial_outputs_and_nonzero_status(self):
        self.services(reports=[self.report, RuntimeError("Final match unavailable")])
        result, _ = self.run_workflow()
        self.assertEqual(result, 1)
        self.assertIn("Final match unavailable", self.diagnostics())
        self.assertEqual(self.docx.read_bytes(), (self.archive() / "initial.docx").read_bytes())
        self.assertEqual(self.pdf.read_bytes(), (self.archive() / "initial.pdf").read_bytes())
        outcome = json.loads((self.archive() / "outcome.json").read_text())
        self.assertNotIn("ADS-OPS-01", outcome["final_selection"]["priority_order"])

    def test_initial_match_failure_still_leaves_outputs_and_diagnostics(self):
        self.services(reports=[RuntimeError("Initial match unavailable")])
        result, _ = self.run_workflow()
        self.assertEqual(result, 1)
        self.assertTrue(self.docx.exists())
        self.assertTrue(self.pdf.exists())
        self.assertIn("Initial match unavailable", self.diagnostics())

    def test_score_decline_does_not_trigger_extra_calls_or_discard_valid_revision(self):
        self.services(reports=[self.report, make_report(weak=False, first_score=7)])
        result, _ = self.run_workflow()
        self.assertEqual(result, 0)
        self.assertEqual(self.select.call_count, 4)
        self.assertEqual(self.match.call_count, 2)
        self.assertIn("Status: revised", self.diagnostics())

    def test_candidate_page_fitting_updates_final_ids_and_diagnostics(self):
        self.services(layout=[SimpleNamespace(has_overflow=False, page_count=2),
                              SimpleNamespace(has_overflow=True, page_count=3),
                              SimpleNamespace(has_overflow=False, page_count=2)])
        result, output = self.run_workflow()
        self.assertEqual(result, 0)
        # The new lowest-priority bullet was removed, restoring the initial rendered content.
        self.assertIn("Status: unchanged_after_fitting", self.diagnostics())
        self.assertNotIn("ADS-OPS-01", output)
        self.assertIn("Removed ADS-OPS-01", self.diagnostics())
        self.assertEqual(self.match.call_count, 1)

    def test_feedback_model_failure_preserves_initial_files(self):
        self.services(responses=[self.initial, self.initial, RuntimeError("Selection API unavailable")])
        result, _ = self.run_workflow()
        self.assertEqual(result, 1)
        self.assertIn("Selection API unavailable", self.diagnostics())
        self.assertEqual(self.docx.read_bytes(), (self.archive() / "initial.docx").read_bytes())
        self.assertEqual(self.pdf.read_bytes(), (self.archive() / "initial.pdf").read_bytes())

    def test_overflow_failure_does_not_publish_or_claim_outputs(self):
        self.services(layout=lambda _: SimpleNamespace(has_overflow=True, page_count=3))
        result, output = self.run_workflow()
        self.assertEqual(result, 1)
        self.assertFalse(self.docx.exists())
        self.assertFalse(self.pdf.exists())
        self.assertIn("No resume was published", output)
        self.assertFalse(json.loads((self.archive() / "outcome.json").read_text())["outputs_published"])
        self.assertLessEqual(self.exports.call_count, len(self.initial.priority_order) + 1)

    def test_partial_publication_restores_both_files(self):
        source_docx, source_pdf = self.root / "new.docx", self.root / "new.pdf"
        target_docx, target_pdf = self.root / "old.docx", self.root / "old.pdf"
        source_docx.write_bytes(b"new docx")
        source_pdf.write_bytes(b"new pdf")
        target_docx.write_bytes(b"old docx")
        target_pdf.write_bytes(b"old pdf")
        from feedback_revision import _atomic_copy
        calls = 0

        def copy(source, destination):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError("Cannot publish PDF")
            _atomic_copy(source, destination)

        with patch("approved_workflow._atomic_copy", side_effect=copy):
            with self.assertRaisesRegex(OSError, "Cannot publish"):
                publish_pair(source_docx=source_docx, source_pdf=source_pdf, docx=target_docx, pdf=target_pdf, archive=self.root)
        self.assertEqual(target_docx.read_bytes(), b"old docx")
        self.assertEqual(target_pdf.read_bytes(), b"old pdf")


if __name__ == "__main__":
    unittest.main()
