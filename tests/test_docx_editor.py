from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from docx_editor import load_resume_template, write_tailored_resume
from resume_models import EvidenceBullet, JobMetadata, ROLE_SPECS, TailoredResume, TailoredRole


class DocxEditorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.template_path = self.root / "stub.docx"
        self.output_path = self.root / "output.docx"
        self._create_template()

    def _create_template(self) -> None:
        document = Document()
        document.add_paragraph("ILAN COOKE")
        document.add_paragraph("[SUMMARY GOES HERE]")
        document.add_paragraph("WORK EXPERIENCE")
        document.add_paragraph("Data Scientist, Marketplace Fraud & Abuse | dates")
        document.add_paragraph("[BULLET POINTS GO HERE]", style="List Bullet")
        document.add_paragraph("Business Intelligence Engineer | dates")
        document.add_paragraph("Business Analyst | dates")
        document.add_paragraph("Related Accounts Enforcement")
        document.add_paragraph("[BULLET POINTS GO HERE]", style="List Bullet")
        document.add_paragraph("Manager, Business Analytics | dates")
        document.add_paragraph("[BULLET POINTS GO HERE]", style="List Bullet")
        document.add_paragraph("Customer Insights Analyst, Marketing Analytics | dates")
        document.add_paragraph("[BULLET POINTS GO HERE]", style="List Bullet")
        document.add_paragraph("EDUCATION")
        document.save(self.template_path)

    def _tailoring(self) -> TailoredResume:
        roles = []
        for spec in ROLE_SPECS:
            roles.append(
                TailoredRole(
                    role_id=spec.role_id,
                    company=spec.company,
                    role=spec.role,
                    bullets=[
                        EvidenceBullet(
                            text=f"{spec.role_id} bullet {index}",
                            source_files=[f"source-{spec.role_id}.md"],
                        )
                        for index in range(1, spec.preferred_bullets + 1)
                    ],
                )
            )
        return TailoredResume(
            summary="Evidence-based professional summary.",
            summary_source_files=["source-amazon_data_scientist.md"],
            roles=roles,
            metadata=JobMetadata(company_name="Example", job_id="123"),
        )

    def test_expands_placeholders_and_preserves_source(self) -> None:
        source_hash = hashlib.sha256(self.template_path.read_bytes()).hexdigest()
        template = load_resume_template(self.template_path)

        result = write_tailored_resume(
            template=template,
            tailoring=self._tailoring(),
            output_path=self.output_path,
        )

        self.assertEqual(result, self.output_path.resolve())
        self.assertEqual(source_hash, hashlib.sha256(self.template_path.read_bytes()).hexdigest())
        output = Document(self.output_path)
        texts = [paragraph.text for paragraph in output.paragraphs]
        self.assertNotIn("[SUMMARY GOES HERE]", texts)
        self.assertNotIn("[BULLET POINTS GO HERE]", texts)
        self.assertIn("Evidence-based professional summary.", texts)
        for spec in ROLE_SPECS:
            generated = [text for text in texts if text.startswith(f"{spec.role_id} bullet")]
            self.assertEqual(len(generated), spec.preferred_bullets)

        generated_paragraphs = [
            paragraph
            for paragraph in output.paragraphs
            if any(
                paragraph.text.startswith(f"{spec.role_id} bullet")
                for spec in ROLE_SPECS
            )
        ]
        self.assertEqual(
            len(generated_paragraphs),
            sum(spec.preferred_bullets for spec in ROLE_SPECS),
        )
        for paragraph in generated_paragraphs:
            num_properties = paragraph._p.find(
                qn("w:pPr") + "/" + qn("w:numPr")
            )
            self.assertIsNotNone(num_properties, paragraph.text)
            self.assertEqual(
                num_properties.find(qn("w:ilvl")).get(qn("w:val")), "0"
            )
            self.assertIsNotNone(num_properties.find(qn("w:numId")))

        descriptor = next(
            paragraph
            for paragraph in output.paragraphs
            if paragraph.text == "Related Accounts Enforcement"
        )
        self.assertIsNone(
            descriptor._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
        )

    def test_allows_fewer_than_preferred_bullets(self) -> None:
        template = load_resume_template(self.template_path)
        tailoring = self._tailoring()
        first_role = tailoring.roles[0]
        tailoring.roles[0] = first_role.model_copy(
            update={"bullets": first_role.bullets[:1]}
        )

        write_tailored_resume(
            template=template,
            tailoring=tailoring,
            output_path=self.output_path,
        )

        output = Document(self.output_path)
        generated = [
            paragraph.text
            for paragraph in output.paragraphs
            if paragraph.text.startswith("amazon_data_scientist bullet")
        ]
        self.assertEqual(len(generated), 1)

    def test_changes_only_document_xml_package_part(self) -> None:
        template = load_resume_template(self.template_path)
        write_tailored_resume(
            template=template,
            tailoring=self._tailoring(),
            output_path=self.output_path,
        )

        with ZipFile(self.template_path) as source, ZipFile(self.output_path) as output:
            self.assertEqual(set(source.namelist()), set(output.namelist()))
            for name in source.namelist():
                if name != "word/document.xml":
                    self.assertEqual(source.read(name), output.read(name), name)


if __name__ == "__main__":
    unittest.main()
